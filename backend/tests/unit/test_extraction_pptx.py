import io

from pptx import Presentation

from app.services.documents.extraction.pptx import extract_pptx


def _make_pptx(slides: list[tuple[str, str | None]]) -> bytes:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    for title, notes in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        if notes is not None:
            slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def test_extracts_slide_text_and_speaker_notes():
    content = _make_pptx(
        [("Q3 Results", "Beat consensus on revenue"), ("Outlook", None)]
    )
    result = extract_pptx(content)

    assert len(result.pages) == 2
    assert result.pages[0].page_number == 1
    assert "Q3 Results" in result.pages[0].text
    assert "Beat consensus on revenue" in result.pages[0].text
    assert result.pages[1].page_number == 2
    assert "Outlook" in result.pages[1].text
    assert result.facts == []


def test_empty_slide_is_flagged_low_text():
    content = _make_pptx([("", None)])
    result = extract_pptx(content)
    assert result.pages[0].quality == "low_text"
