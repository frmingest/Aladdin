"""Small synthetic-file builders shared by unit and integration tests — no
fixture binaries checked into the repo, everything is generated in-memory."""

import io


DEFAULT_CSV_HEADER = "Ticker,Name,Asset class,Quantity,Weight %,Cost basis,Currency,Sector/Theme,Notes"


def make_portfolio_csv(rows: list[str], header: str = DEFAULT_CSV_HEADER) -> bytes:
    return ("\n".join([header, *rows]) + "\n").encode("utf-8")


def make_xlsx(sheet_rows: dict[str, list[list]]) -> bytes:
    """sheet_rows: {sheet_name: [[row1 cells], [row2 cells], ...]} — first row
    of each sheet is treated as the header/period row by the app's extractor."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheet_rows.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def make_pptx(slides: list[tuple[str, str | None]]) -> bytes:
    """slides: [(slide_text, speaker_notes_or_None), ...]."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[1]
    for text, notes in slides:
        slide = prs.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        box.text_frame.text = text
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def make_pdf(pages: list[str]) -> bytes:
    import fitz

    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
    content = doc.tobytes()
    doc.close()
    return content
