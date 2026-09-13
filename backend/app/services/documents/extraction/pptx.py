"""PowerPoint extraction (§6.2): slide text + speaker notes, slide numbers
retained as page_number. Chart-heavy-slide rendering to images (for vision
input) is deferred to Phase 3, alongside the rest of the LLM pipeline."""

import io

from app.services.documents.extraction.base import ExtractedPage, ExtractionResult, quality_for_text


def extract_pptx(content: bytes) -> ExtractionResult:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    pages: list[ExtractedPage] = []

    for i, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                parts.append(f"[Speaker notes]\n{notes_text}")
        text = "\n\n".join(parts)
        pages.append(ExtractedPage(page_number=i, text=text, quality=quality_for_text(text)))

    return ExtractionResult(pages=pages, facts=[], quality_flags=[])
